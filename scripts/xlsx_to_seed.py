#!/usr/bin/env python3
"""
Convert catfood.xlsx (Tom's research spreadsheet) into database/seeds/catfood_seed.json.

Usage:  python3 scripts/xlsx_to_seed.py path/to/catfood.xlsx [out.json]

Every product keeps where it came from (sheet + row), the spreadsheet's own link, and a list of data "flags"
(missing values, #DIV/0!, unparseable cells) so the later audit knows what to look at. Nothing here is a label value:
numbers are whatever the spreadsheet said. No barcodes exist in the spreadsheet; products are imported without a GTIN
and get one when scanned (Scan -> "pick the existing product").

Layout assumptions (checked against the Dec 2024 file):
  wet sheets  : Name | text columns... | Primary ingredients | then 8 numbers
                [packs, can_oz, case_oz, case_price, price_per_oz, kcal_per_can, kcal_per_oz, price_per_kcal]
                then Protein, Fat, Carb, Fiber, Moisture, Link (the header row has one fewer numeric label than the data)
  Dry Food    : headers map 1:1 (Size (lb), Cups, Price, Price/lb, $/100cal, Protein, Fat, Fiber, Carb, Moisture, kcals, Link)
"""
import json
import re
import sys
from datetime import datetime, timezone

import openpyxl

SOURCE = "catfood.xlsx"
SKIP_NAMES = {"Color Key"}
# Sheets whose first four nutrition columns are already percentages (75.4) rather than fractions (0.754)
PERCENT_SHEETS = {"Weruva"}
PRICE_DATE = "2024-12"  # file last saved Dec 2024; prices are a snapshot, never current

BRAND_BY_SHEET = {
    "Tiki Cat": "Tiki Cat", "Weruva": "Weruva", "Purina": "Purina", "Nacho": "Nacho",
    "Instinct": "Instinct", "Hills": "Hill's", "Fancy Feast": "Fancy Feast",
}


def num(v):
    """float, or None for blanks / Excel errors / junk. Second value: the raw text if it was junk."""
    if v is None or v == "":
        return None, None
    if isinstance(v, (int, float)):
        return float(v), None
    s = str(v).strip()
    if s.startswith("#") or s == "":
        return None, s or None
    try:
        return float(s.rstrip("%")), None
    except ValueError:
        return None, s


def pct(v, flags, label, as_is=False):
    """Percent value as a number (40.6). The sheet mixes fractions (0.406) and percents (75.4)."""
    n, raw = num(v)
    if raw:
        flags.append(f"{label} unreadable in sheet: {raw!r}")
    if n is None:
        return None
    return round(n if as_is or n > 1 else n * 100, 1)


def text(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def clean_link(v):
    s = text(v)
    if not s:
        return None
    if not s.startswith("http"):
        s = "https://" + s
    return s


def nonempty(row):
    return [(i, c) for i, c in enumerate(row) if c not in (None, "")]


def wet_sheet(ws, sheet):
    brand = BRAND_BY_SHEET[sheet]
    out, line, header = [], None, None
    rows = list(ws.iter_rows(values_only=True))
    for idx, row in enumerate(rows, start=1):
        cells = nonempty(row)
        if not cells:
            continue
        first = text(row[0])
        if first == "Name" and "Protein" in [text(c) for c in row]:
            heads = [text(c) for c in row]
            ing = next((i for i, h in enumerate(heads) if h and h.lower().startswith("primary")), None)
            if ing is None:
                header = None
                continue
            header = {"ing": ing, "heads": heads}
            continue
        # a lone text cell directly above a header row is a product-line label
        if len(cells) == 1 and isinstance(cells[0][1], str):
            nxt = rows[idx] if idx < len(rows) else ()
            if text(nxt[0] if nxt else None) == "Name":
                line = cells[0][1].replace('"', "").strip()
                if line == "Nutritional Example":
                    line = None
            continue
        if not header or not first or first in SKIP_NAMES or first.startswith("Ideal Cat Food"):
            continue

        ing_i = header["ing"]
        heads = header["heads"]
        flags = []
        meta = {"sheet": sheet, "row": idx}

        # Text columns between Name and Primary ingredients, labelled by their header
        texture = variety = None
        for i in range(1, ing_i):
            label, val = heads[i], text(row[i])
            if not label or not val:
                continue
            key = label.lower()
            if key == "texture":
                texture = val
            elif key in ("variety", "flavor"):
                variety = val
                meta[key] = val
            elif key == "protein":
                meta["protein_type"] = val

        nums = [row[ing_i + 1 + k] if ing_i + 1 + k < len(row) else None for k in range(14)]
        packs, can_oz, case_oz, price, ppo, kcal_can, kcal_oz, ppk = (num(x)[0] for x in nums[:8])
        for label, v in zip(["packs", "can_oz", "case_oz", "case_price", "price_per_oz", "kcal_per_can", "kcal_per_oz", "price_per_kcal"], nums[:8]):
            _, raw = num(v)
            if raw:
                flags.append(f"{label} unreadable in sheet: {raw!r}")

        # If the sheet only has a first-link-less row (Excel-error-only rows like Fancy Feast "Purely"), nutrition is absent.
        nutrition = {
            "protein_pct": pct(nums[8], flags, "protein", sheet in PERCENT_SHEETS),
            "fat_pct": pct(nums[9], flags, "fat", sheet in PERCENT_SHEETS),
            "carb_pct": pct(nums[10], flags, "carb", sheet in PERCENT_SHEETS),
            "fiber_pct": pct(nums[11], flags, "fiber", sheet in PERCENT_SHEETS),
            "moisture_pct": pct(nums[12], flags, "moisture"),
        }
        if kcal_can:
            nutrition["kcal_per_can"] = kcal_can
        if kcal_oz:
            nutrition["kcal_per_oz"] = round(kcal_oz, 2)
        nutrition = {k: v for k, v in nutrition.items() if v is not None}
        if nutrition:
            nutrition["basis"] = "spreadsheet value, not from a label; probably dry-matter. Audit before trusting."

        # The sheet has a few more cells after moisture: product link(s), occasional scribbles
        extras = [text(c) for c in row[ing_i + 14:] if text(c)]
        links = [clean_link(x) for x in extras if re.search(r"\.(com|net|org)", x, re.I) or x.startswith("http")]
        scribbles = [x for x in extras if not (re.search(r"\.(com|net|org)", x, re.I) or x.startswith("http"))]
        if scribbles:
            meta["sheet_scribble"] = scribbles

        ingredients = text(row[ing_i])
        if not ingredients:
            flags.append("no ingredients in sheet")
        if not nutrition:
            flags.append("no nutrition values in sheet")
        if not links:
            flags.append("no product link in sheet")

        meta.update({
            "packs_per_case": packs, "can_oz": can_oz, "case_oz": case_oz, "case_price": price,
            "price_per_oz": ppo, "price_snapshot": PRICE_DATE, "flags": flags,
        })
        meta = {k: v for k, v in meta.items() if v not in (None, [], {})}

        out.append({
            "brand": brand, "name": first, "line": line, "texture": texture, "form": "wet",
            "ingredients": ingredients, "nutrition": nutrition or None,
            "source_url": links[0] if links else None, "meta": {**meta, **({"extra_links": links[1:]} if len(links) > 1 else {})},
            "_variety": variety,
        })
    return out


def dry_sheet(ws):
    out, line = [], None
    rows = list(ws.iter_rows(values_only=True))
    heads = None
    for idx, row in enumerate(rows, start=1):
        cells = nonempty(row)
        if not cells:
            continue
        first = text(row[0])
        if first == "Name":
            heads = [text(c) for c in row]
            continue
        if len(cells) == 1 and isinstance(cells[0][1], str):
            line = cells[0][1].strip()
            continue
        if not heads or not first:
            continue
        d = {h: row[i] for i, h in enumerate(heads) if h and i < len(row)}
        if not isinstance(d.get("Protein"), (int, float)):
            continue  # the "Ideal Cat Food Balance Example" row
        flags = []
        nutrition = {
            "protein_pct": pct(d.get("Protein"), flags, "protein"), "fat_pct": pct(d.get("Fat"), flags, "fat"),
            "fiber_pct": pct(d.get("Fiber"), flags, "fiber"), "carb_pct": pct(d.get("Carb"), flags, "carb"),
            "moisture_pct": pct(d.get("Moisture"), flags, "moisture"),
        }
        kcal, _ = num(d.get("kcals"))
        if kcal:
            nutrition["kcal_per_cup"] = kcal
        nutrition = {k: v for k, v in nutrition.items() if v is not None}
        if nutrition:
            nutrition["basis"] = "spreadsheet value, not from a label; probably dry-matter. Audit before trusting."
        brand, _, name = first.partition(" - ")
        if not name:
            m = re.match(r"^(Dr\.\s+\S+|\S+)\s*(.*)$", first)
            brand, name = m.group(1), (m.group(2) or first)
        link = clean_link(d.get("Link"))
        if not link:
            flags.append("no product link in sheet")
        meta = {"sheet": "Dry Food", "row": idx, "cost_tier": line, "size_lb": num(d.get("Size (lb)"))[0],
                "cups": num(d.get("Cups"))[0], "price": num(d.get("Price"))[0], "price_snapshot": PRICE_DATE, "flags": flags,
                "brand_guessed_from_name": True}
        out.append({
            "brand": brand.strip(), "name": name.strip(), "line": None, "texture": None, "form": "dry",
            "ingredients": text(d.get("Primary ingredients")), "nutrition": nutrition or None, "source_url": link,
            "meta": {k: v for k, v in meta.items() if v not in (None, [], {})}, "_variety": None,
        })
    return out


def main():
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else "database/seeds/catfood_seed.json"
    wb = openpyxl.load_workbook(src, data_only=True)
    products = []
    for ws in wb:
        products += dry_sheet(ws) if ws.title == "Dry Food" else wet_sheet(ws, ws.title)

    seen = {}
    for p in products:
        variety = p.pop("_variety")
        key = "|".join([SOURCE, p["meta"]["sheet"], p["line"] or "", variety or "", p["name"], p["texture"] or ""]).lower()
        seen[key] = seen.get(key, 0) + 1
        if seen[key] > 1:
            key += f"#{seen[key]}"
            p["meta"].setdefault("flags", []).append("same name/line appears more than once in the sheet")
        p["import_key"] = key
        p.update({"species": "cat", "kind": "food", "source": SOURCE, "audit_status": "unreviewed"})
        if variety:
            p["name"] = f"{p['name']} ({variety})" if p["meta"].get("variety") == variety and p["brand"] != "Weruva" else p["name"]
        notes = list(p["meta"].get("flags", []))
        if p["meta"].get("brand_guessed_from_name"):
            notes.append("brand/name split guessed from one sheet cell")
        p["audit_notes"] = "; ".join(notes) or None

    doc = {"source": SOURCE, "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "products": products}
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)
    flagged = sum(1 for p in products if p["meta"].get("flags"))
    print(f"{len(products)} products -> {dst} ({flagged} with data flags)")


if __name__ == "__main__":
    main()
